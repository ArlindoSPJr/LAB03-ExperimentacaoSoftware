"""Enriquece a planilha de rotulagem com o conteúdo público do GitHub (apoio à S02-O2.3).

Para cada repositório: descrição, tópicos, linguagem e o README (resumido).
Para cada release: as notas da release e os títulos dos commits (`compare` com a
release anterior). Nada aqui sugere rótulo: a planilha **não** recebe nenhuma saída
da heurística, e as colunas de rótulo (`tipo_projeto`, `entregas_reais`,
`corretiva`, `observacao`) são copiadas da entrada sem alteração.

Uso:
  python -m validacao.enriquecer_planilha --entrada rotulos/B.csv --saida rotulos/B_preencher.xlsx
  (aceita .csv ou .xlsx na entrada; token em GITHUB_TOKEN; respostas ficam no cache)

Requer `openpyxl` (só para gravar o .xlsx): pip install openpyxl
"""
from __future__ import annotations

import argparse
import base64
import os
import pathlib
import re
import sys
from urllib.parse import unquote

import pandas as pd

from pipeline.cache import DiskCache
from pipeline.http import GitHubClient, NotFoundError
from pipeline.releases import fetch_commits
from validacao.concordancia import DIMENSOES

ROTULOS = ["tipo_projeto", "entregas_reais", "corretiva", "observacao"]
MAX_README = 2500
MAX_NOTAS = 1500
MAX_COMMITS = 25


def _limpar(texto: str | None, limite: int) -> str:
    texto = re.sub(r"<!--.*?-->", "", texto or "", flags=re.S)
    texto = re.sub(r"<img[^>]*>|!\[[^\]]*\]\([^)]*\)", "", texto)  # imagens e badges
    texto = re.sub(r"\n\s*\n+", "\n", texto).strip()
    return texto[:limite] + " […]" if len(texto) > limite else texto


def info_repo(client, repo: str) -> tuple[str, str]:
    """(descrição + tópicos + linguagem, README resumido)."""
    try:
        meta, _ = client.get(f"/repos/{repo}")
    except NotFoundError:
        return "(repositório não encontrado)", ""
    topicos = ", ".join(meta.get("topics") or [])
    descricao = meta.get("description") or "(sem descrição)"
    if topicos:
        descricao += f"\nTópicos: {topicos}"
    descricao += f"\nLinguagem: {meta.get('language')}"
    try:
        readme, _ = client.get(f"/repos/{repo}/readme")
        texto = base64.b64decode(readme.get("content", "")).decode("utf-8", errors="replace")
    except NotFoundError:
        texto = "(sem README)"
    return descricao, _limpar(texto, MAX_README)


def notas_release(client, repo: str, tag: str) -> str:
    try:
        rel, _ = client.get(f"/repos/{repo}/releases/tags/{tag}")
    except NotFoundError:
        return "(release não encontrada)"
    titulo = rel.get("name") or ""
    cabecalho = f"[{titulo}]\n" if titulo and titulo != tag else ""
    return _limpar(cabecalho + (rel.get("body") or "(notas vazias)"), MAX_NOTAS)


def commits_release(client, repo: str, url_commits: str) -> str:
    m = re.search(r"/compare/(.+)\.\.\.(.+)$", url_commits or "")
    if not m:
        return "(primeira release: sem compare)"
    commits = fetch_commits(client, repo, unquote(m.group(1)), unquote(m.group(2)))
    if commits is None:
        return "(compare indisponível)"
    titulos = [c["message"].splitlines()[0][:110] for c in commits]
    texto = f"{len(titulos)} commits:\n- " + "\n- ".join(titulos[:MAX_COMMITS])
    if len(titulos) > MAX_COMMITS:
        texto += f"\n… (+{len(titulos) - MAX_COMMITS})"
    return texto


def enriquecer(client, planilha: pd.DataFrame) -> pd.DataFrame:
    df = planilha.copy()
    for col in ROTULOS:
        if col not in df:
            df[col] = ""
    readme_por_repo: dict[str, tuple[str, str]] = {}
    desc, readme, notas, commits = [], [], [], []
    for row in df.itertuples(index=False):
        if row.nivel == "repo":
            readme_por_repo[row.repo] = info_repo(client, row.repo)
            d, r = readme_por_repo[row.repo]
            desc.append(d), readme.append(r), notas.append(""), commits.append("")
        else:
            desc.append(""), readme.append("")
            notas.append(notas_release(client, row.repo, row.tag))
            commits.append(commits_release(client, row.repo, row.url_commits))
    pos = list(df.columns).index("url") + 1
    for i, (nome, valores) in enumerate([("descricao_repo", desc), ("readme", readme),
                                         ("notas_release", notas), ("commits_release", commits)]):
        if nome in df:
            df = df.drop(columns=nome)
        df.insert(pos + i, nome, valores)
    return df


def gravar_xlsx(df: pd.DataFrame, caminho: pathlib.Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "rotulagem"
    cols = list(df.columns)
    idx = {c: i + 1 for i, c in enumerate(cols)}
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    azul, amarelo, cinza = (PatternFill("solid", fgColor=c) for c in ("D9E1F2", "FFF2CC", "EDEDED"))
    listas = {}
    for dim, valores in DIMENSOES.items():
        listas[dim] = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True,
                                     showErrorMessage=True, error="Escolha um valor da lista.")
        ws.add_data_validation(listas[dim])
    quebra = Alignment(wrap_text=True, vertical="top")
    for r, row in enumerate(df.itertuples(index=False), start=2):
        ws.append(["" if pd.isna(v) else v for v in row])
        eh_repo = row.nivel == "repo"
        for c in cols:
            ws.cell(r, idx[c]).alignment = quebra
        if eh_repo:
            for c in ("nivel", "repo", "descricao_repo", "readme"):
                ws.cell(r, idx[c]).fill = azul
        for c in ("url", "url_commits"):
            cell = ws.cell(r, idx[c])
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
        preencher = ["tipo_projeto", "entregas_reais"] if eh_repo else ["corretiva"]
        bloquear = ["corretiva"] if eh_repo else ["tipo_projeto", "entregas_reais"]
        for c in preencher:
            ws.cell(r, idx[c]).fill = amarelo
            listas[c].add(ws.cell(r, idx[c]))
        for c in bloquear:
            ws.cell(r, idx[c]).fill = cinza
        ws.row_dimensions[r].height = 150
    larguras = {"nivel": 7, "repo": 22, "tag": 12, "published_at": 11, "url": 14, "descricao_repo": 35,
                "readme": 60, "notas_release": 70, "commits_release": 55, "url_commits": 14,
                "tipo_projeto": 20, "entregas_reais": 13, "corretiva": 10, "observacao": 30}
    for c, w in larguras.items():
        if c in idx:
            ws.column_dimensions[get_column_letter(idx[c])].width = w
    ws.freeze_panes = "D2"
    wb.save(caminho)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m validacao.enriquecer_planilha", description=__doc__.splitlines()[0])
    parser.add_argument("--entrada", required=True, help="planilha de rotulagem (.csv ou .xlsx)")
    parser.add_argument("--saida", required=True, help="arquivo .xlsx de saída")
    parser.add_argument("--cache", default="cache", help="diretório do cache (padrão: cache)")
    args = parser.parse_args(argv)
    token = os.environ.get("GITHUB_TOKEN")
    if not token:
        print("erro: defina a variável de ambiente GITHUB_TOKEN.", file=sys.stderr)
        return 1
    entrada = pathlib.Path(args.entrada)
    ler = pd.read_excel if entrada.suffix.lower() == ".xlsx" else pd.read_csv
    planilha = ler(entrada, dtype=str).fillna("")
    client = GitHubClient(token, DiskCache(pathlib.Path(args.cache)))
    gravar_xlsx(enriquecer(client, planilha), pathlib.Path(args.saida))
    print(f"planilha enriquecida em {args.saida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
