"""Pré-rotulagem automática (assistida por IA) da planilha de rotulagem da S02.

Lê a planilha enriquecida (`descricao_repo`, `readme`, `notas_release`, `commits_release`)
e sugere `tipo_projeto`, `entregas_reais` e `corretiva` por regras de palavras-chave.
Só preenche quando a evidência é clara; os casos duvidosos ficam **em branco**, com
`revisar = sim` e a sugestão/motivo ao lado, para decisão humana. Rótulos já
preenchidos à mão nunca são sobrescritos. A coluna `origem_rotulo` registra se o rótulo
é `manual`, `automatico` ou `pendente`, para a Metodologia declarar o processo.

Uso:
  python -m validacao.pre_rotular --entrada B_rotulagem_com_readme.xlsx --saida rotulos/B_pre_rotulado.xlsx
"""
from __future__ import annotations

import argparse
import pathlib
import re
import sys

import pandas as pd

from validacao.enriquecer_planilha import gravar_xlsx

COLS_AUTO = ["origem_rotulo", "revisar", "sugestao_auto", "motivo_auto"]

# --- tipo do projeto -------------------------------------------------------------------
# Palavras genéricas (server, service, application, package, plugin, terminal) ficam de fora:
# aparecem tanto em bibliotecas quanto em aplicações e geravam rótulos errados (aiohttp, tabby).
SINAIS_TIPO = {
    "ferramenta CLI": [r"command[- ]line", r"\bcli\b", r"brew install", r"cargo install", r"go install",
                       r"pipx install", r"npm (?:i|install) -g", r"\busage:"],
    "biblioteca/framework": [r"\blibrary\b", r"\bframework\b", r"\bsdk\b", r"pip install",
                             r"npm (?:i|install) (?!-g)", r"\bcrate\b", r"\bbindings?\b", r"\bimport \w",
                             r"\brequire\(", r"\bapi client\b", r"\btoolkit\b"],
    "aplicacao/servico": [r"self[- ]?hosted", r"\bweb ?app\b", r"\bdashboard\b", r"docker[- ]compose",
                          r"docker run", r"\bdesktop\b", r"\bandroid\b", r"\bios\b", r"\bwiki\b"],
}
SINAIS_OUTRO = [r"\bawesome\b", r"\bdataset\b", r"\bcurated list\b", r"\bdocumentation site\b", r"\bdotfiles\b"]


def _contar(padroes: list[str], texto: str) -> int:
    return sum(len(re.findall(p, texto)) for p in padroes)


def classificar_tipo(descricao: str, readme: str) -> tuple[str | None, str, str]:
    """(rótulo ou None se duvidoso, sugestão, motivo). Descrição/tópicos valem o triplo do README."""
    desc, corpo = (descricao or "").lower(), (readme or "").lower()[:2500]
    pontos = {t: 3 * _contar(p, desc) + min(_contar(p, corpo), 4) for t, p in SINAIS_TIPO.items()}
    ordem = sorted(pontos.items(), key=lambda kv: -kv[1])
    (top, p1), (_, p2) = ordem[0], ordem[1]
    motivo = ", ".join(f"{t}={p}" for t, p in ordem)
    if _contar(SINAIS_OUTRO, desc) and p1 < 4:
        return None, "outro", f"indícios de 'outro' ({motivo})"
    if p1 >= 4 and p1 >= 2.5 * p2:
        return top, top, motivo
    return None, top if p1 else "", f"evidência fraca ou empatada ({motivo})"


# --- releases são entregas reais? --------------------------------------------------------
TAG_NAO_ENTREGA = re.compile(r"nightly|snapshot|canary|\bdev\b|-dev|build[-_.]?\d|^\d{8,}|latest|ci-", re.I)


def _notas_vazias(notas: str) -> bool:
    n = (notas or "").strip()
    return not n or n.startswith("(") or len(re.sub(r"\[.*?\]", "", n).strip()) < 20


def classificar_entregas(tags: list[str], notas: list[str]) -> tuple[str | None, str, str]:
    n = len(tags) or 1
    automaticas = sum(bool(TAG_NAO_ENTREGA.search(t or "")) for t in tags)
    vazias = sum(_notas_vazias(x) for x in notas)
    motivo = f"{n - vazias}/{n} releases com notas; {automaticas}/{n} tags tipo nightly/build"
    if automaticas == 0 and vazias <= 1:
        return "sim", "sim", motivo
    if automaticas >= 3:
        return None, "nao", motivo
    return None, "incerto", motivo


# --- release corretiva? ------------------------------------------------------------------
FIX = re.compile(r"\b(fix(e[sd])?|bug ?fix(es)?|hotfix|revert|regression|crash|patch(ed)?|security|cve-\d|"
                 r"resolve[sd]? (an? )?(issue|bug|error)|correct(ed|s)?)\b", re.I)
FEAT = re.compile(r"\b(feat|features?|add(ed|s)?|new|introduc\w*|implement\w*|support for|enhance\w*|"
                  r"improve\w*|refactor\w*)\b", re.I)
NEUTRO = re.compile(r"\b(bump|chore\(deps|dependabot|update dependenc|renovate|merge (branch|pull)|"
                    r"release v?\d|version bump|changelog)\b"
                    r"|^[-*•]\s*(ci|build|chore|docs|test|style)(\(.*?\))?!?:", re.I)
SECAO_FIX = re.compile(r"^#+.*\b(bug ?fix|fix|patch|security)", re.I)
SECAO_FEAT = re.compile(r"^#+.*(\b(feature|new|added|enhancement|improvement|highlight|major|minor|breaking)"
                        r"|新功能|功能|改进)", re.I)
# x.y ou x.y.0: bump major/minor costuma trazer novidade; correção pura aí vai para revisão
TAG_NAO_PATCH = re.compile(r"(?:^|[^\d.])\d+\.\d+(?:\.0)?$")


def _contar_linhas(texto: str) -> tuple[int, int]:
    fix = feat = 0
    secao = None
    for linha in (texto or "").splitlines():
        s = linha.strip()
        if s.startswith("#"):
            secao = "fix" if SECAO_FIX.search(s) else "feat" if SECAO_FEAT.search(s) else None
            continue
        if not re.match(r"^([-*•]|\d+\.)\s", s) or NEUTRO.search(s):
            continue
        if secao == "fix" or (secao is None and FIX.search(s) and not FEAT.match(s.lstrip("-*• "))):
            fix += 1
        elif secao == "feat" or FEAT.search(s):
            feat += 1
    return fix, feat


def classificar_corretiva(notas: str, commits: str, tag: str = "") -> tuple[str | None, str, str]:
    f1, a1 = _contar_linhas(notas)
    f2, a2 = _contar_linhas(commits)
    fix, feat = f1 + f2, a1 + a2
    motivo = f"itens de correção={fix}, de funcionalidade={feat} (notas {f1}/{a1}, commits {f2}/{a2})"
    if fix + feat == 0:
        return None, "", "sem itens classificáveis nas notas/commits"
    razao = fix / (fix + feat)
    sugestao = "sim" if razao >= 0.5 else "nao"
    secao_feat = any(SECAO_FEAT.search(linha.strip()) for linha in (notas or "").splitlines())
    if fix >= 3 and razao >= 0.8:
        if secao_feat or TAG_NAO_PATCH.search(tag or ""):
            return None, "sim", f"correções, mas notas com novidades ou bump major/minor: {motivo}"
        return "sim", "sim", motivo
    if feat >= 3 and razao <= 0.2:
        return "nao", "nao", motivo
    return None, sugestao, f"misto: {motivo}"


# --- orquestração ------------------------------------------------------------------------
def _vazio(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan")


def pre_rotular(planilha: pd.DataFrame) -> pd.DataFrame:
    df = planilha.copy()
    for c in COLS_AUTO:
        df[c] = ""
    for repo, grupo in df.groupby("repo", sort=False):
        rels = grupo[grupo["nivel"] == "release"]
        for i, row in grupo.iterrows():
            if row["nivel"] == "repo":
                tipo = classificar_tipo(row.get("descricao_repo"), row.get("readme"))
                entregas = classificar_entregas(rels["tag"].fillna("").tolist(),
                                                rels["notas_release"].fillna("").tolist())
                alvos = [("tipo_projeto", tipo), ("entregas_reais", entregas)]
            else:
                alvos = [("corretiva", classificar_corretiva(row.get("notas_release"), row.get("commits_release"),
                                                           row.get("tag")))]
            origens, sugestoes, motivos = [], [], []
            for col, (rotulo, sugestao, motivo) in alvos:
                if not _vazio(row[col]):
                    origens.append("manual")
                elif rotulo is not None:
                    df.at[i, col] = rotulo
                    origens.append("automatico")
                else:
                    origens.append("pendente")
                    sugestoes.append(f"{col}={sugestao or '?'}")
                    motivos.append(f"{col}: {motivo}")
                    continue
                if origens[-1] == "automatico":
                    motivos.append(f"{col}: {motivo}")
            df.at[i, "origem_rotulo"] = "/".join(origens)
            df.at[i, "revisar"] = "sim" if "pendente" in origens else "nao"
            df.at[i, "sugestao_auto"] = "; ".join(sugestoes)
            df.at[i, "motivo_auto"] = " | ".join(motivos)
    return df


def gravar(df: pd.DataFrame, caminho: pathlib.Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    gravar_xlsx(df, caminho)
    wb = load_workbook(caminho)
    ws = wb["rotulagem"]
    idx = {c.value: c.column for c in ws[1]}
    laranja = PatternFill("solid", fgColor="F8CBAD")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, idx["revisar"]).value == "sim":
            for c in ("revisar", "sugestao_auto"):
                ws.cell(r, idx[c]).fill = laranja
    from openpyxl.utils import get_column_letter
    for c, w in {"origem_rotulo": 16, "revisar": 8, "sugestao_auto": 24, "motivo_auto": 45}.items():
        ws.column_dimensions[get_column_letter(idx[c])].width = w
    ws.auto_filter.ref = ws.dimensions

    resumo = wb.create_sheet("resumo")
    resumo.append(["dimensao", "manual", "automatico", "pendente (revisar)"])
    for cell in resumo[1]:
        cell.font = Font(bold=True)
    for k, col in enumerate(["tipo_projeto", "entregas_reais", "corretiva"]):
        nivel = "repo" if col != "corretiva" else "release"
        origens = df.loc[df["nivel"] == nivel, "origem_rotulo"].str.split("/")
        pos = k if nivel == "repo" else 0
        valores = origens.str[pos].value_counts()
        resumo.append([col] + [int(valores.get(o, 0)) for o in ("manual", "automatico", "pendente")])
    resumo.append([])
    resumo.append(["Rótulos 'automatico' vieram de regras de palavras-chave (validacao/pre_rotular.py) "
                   "e devem ser conferidos; declarar o uso na Metodologia."])
    resumo.column_dimensions["A"].width = 18
    wb.save(caminho)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m validacao.pre_rotular", description=__doc__.splitlines()[0])
    parser.add_argument("--entrada", required=True, help="planilha enriquecida (.xlsx ou .csv)")
    parser.add_argument("--saida", required=True, help="novo .xlsx (a entrada não é alterada)")
    args = parser.parse_args(argv)
    entrada, saida = pathlib.Path(args.entrada), pathlib.Path(args.saida)
    if entrada.resolve() == saida.resolve():
        print("erro: a saída deve ser um arquivo novo.", file=sys.stderr)
        return 1
    ler = pd.read_excel if entrada.suffix.lower() == ".xlsx" else pd.read_csv
    df = pre_rotular(ler(entrada, dtype=str).fillna(""))
    gravar(df, saida)
    print(f"pré-rotulagem em {saida}; {int((df['revisar'] == 'sim').sum())} linhas para revisar")
    return 0


if __name__ == "__main__":
    sys.exit(main())
